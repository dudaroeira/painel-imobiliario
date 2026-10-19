"""ABECIP — financiamento SBPE (unidades e R$, construção × aquisição) e IGMI-R (preço de transação).

Os nomes dos arquivos mudam a cada mês; o coletor raspa a página do indicador para achar o
link `download?file=...xlsx` e cai para um nome conhecido se a página mudar.
"""
from __future__ import annotations

import io
import re

import openpyxl
from bs4 import BeautifulSoup

from .. import http, store

PAGES = {
    "financiamento": ("https://www.abecip.org.br/credito-imobiliario/indicadores/financiamento", "unidades-site81.xlsx"),
    "igmi": ("https://www.abecip.org.br/igmi-r-abecip/serie-historica", "igmi-r-serie-historica-agosto2026.xlsx"),
}
DL = "https://www.abecip.org.br/download?file={name}"


def _find_xlsx(page_url: str, fallback: str, refresh: bool) -> str:
    try:
        html = http.get(page_url, cache_name=f"abecip/{page_url.rstrip('/').split('/')[-1]}.html", refresh=refresh).decode("utf-8", "ignore")
        soup = BeautifulSoup(html, "lxml")
        for a in soup.select("a[href*='download?file=']"):
            name = a["href"].split("file=")[-1]
            if name.lower().endswith(".xlsx"):
                return name
    except Exception as e:
        print(f"[abecip] não achei link em {page_url}: {e}")
    return fallback


def _month(v) -> str | None:
    if hasattr(v, "year"):
        return v.strftime("%Y-%m-01")
    m = re.match(r"^(\d{4})\s+(\d{2})$", str(v).strip())  # IGMI usa '2014 01'
    return f"{m.group(1)}-{m.group(2)}-01" if m else None


def collect(refresh: bool = True):
    out = []
    # --- Financiamento SBPE (aba BD_Unidades: Período | Unid. constr | Unid. aquis | Unid. total | R$ constr | R$ aquis | R$ total)
    name = _find_xlsx(*PAGES["financiamento"], refresh)
    wb = openpyxl.load_workbook(io.BytesIO(http.get(DL.format(name=name), cache_name=f"abecip/{name}", refresh=refresh)),
                                data_only=True, read_only=True)
    ws = wb["BD_Unidades"] if "BD_Unidades" in wb.sheetnames else wb.worksheets[0]
    cols = [(1, "sbpe_unid_construcao", "SBPE — unidades financiadas (construção)", "unidades"),
            (2, "sbpe_unid_aquisicao", "SBPE — unidades financiadas (aquisição)", "unidades"),
            (3, "sbpe_unid_total", "SBPE — unidades financiadas (total)", "unidades"),
            (4, "sbpe_rs_construcao", "SBPE — valor financiado (construção)", "R$ mi"),
            (5, "sbpe_rs_aquisicao", "SBPE — valor financiado (aquisição)", "R$ mi"),
            (6, "sbpe_rs_total", "SBPE — valor financiado (total)", "R$ mi")]
    for r in ws.iter_rows(min_row=6, values_only=True):
        d = _month(r[0]) if r and r[0] is not None else None
        if not d:
            continue
        for c, sid, nm, unit in cols:
            if isinstance(r[c], (int, float)):
                out.append(dict(source="abecip", series_id=sid, name=nm, block="credito", region="BR", unit=unit, freq="M", date=d, value=r[c]))
    # --- IGMI-R (Mês | São Paulo | Rio | BH | Fortaleza | Recife | Curitiba | POA | Salvador | Goiânia | Brasília | BRASIL)
    name = _find_xlsx(*PAGES["igmi"], refresh)
    wb = openpyxl.load_workbook(io.BytesIO(http.get(DL.format(name=name), cache_name=f"abecip/{name}", refresh=refresh)),
                                data_only=True, read_only=True)
    ws = wb.worksheets[0]
    header = None
    for r in ws.iter_rows(values_only=True):
        if header is None:
            if r and r[0] == "Mês":
                header = [("BR" if str(h).upper() == "BRASIL" else h) for h in r]
            continue
        d = _month(r[0]) if r and r[0] is not None else None
        if not d:
            continue
        for i, v in enumerate(r[1:], start=1):
            if i < len(header) and header[i] and isinstance(v, (int, float)):
                out.append(dict(source="abecip", series_id="igmi_r", name="IGMI-R — índice de preços de transação (jan/2014=100)",
                                block="precos", region=header[i], unit="índice", freq="M", date=d, value=v))
    df = store.frame(out)
    store.save(df, "abecip", {"url": PAGES["financiamento"][0], "title": "ABECIP — SBPE e IGMI-R",
                              "coverage": "Brasil (SBPE); Brasil + 10 capitais (IGMI-R, preços de TRANSAÇÃO)"})
    return df


if __name__ == "__main__":
    collect()
